from __future__ import annotations

import json
import html
import re
import os
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from src.core.consolidator import render_separator
from src.core.models import Block, ConsolidationFile, ConsolidationOptions


Writer = Callable[[Path, Sequence[ConsolidationFile], ConsolidationOptions], None]


class WriterRegistry:
    def __init__(self) -> None:
        self._writers: dict[str, Writer] = {}

    def register(self, format: str, writer: Writer) -> None:
        self._writers[format.lower().lstrip(".")] = writer

    def supported_formats(self) -> tuple[str, ...]:
        return tuple(sorted(self._writers))

    def get_writer(self, format: str) -> Writer:
        key = format.lower().lstrip(".")
        if key not in self._writers:
            raise ValueError(f"Unsupported output format: {format}")
        return self._writers[key]


def _write_txt(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        for index, item in enumerate(files):
            if index:
                stream.write(render_separator(options, item))
            if options.separator_mode == "filename":
                stream.write(f"=== {item.name} ===\n")
            stream.write(item.text)


def _json_value(value: object) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"Cannot serialize {type(value).__name__}")


def _record(item: ConsolidationFile) -> dict:
    if item.document is not None:
        return asdict(item.document)
    return {
        "source": {"path": str(item.path), "sha256": None},
        "blocks": [asdict(Block("paragraph", item.text))],
        "metadata": {},
        "warnings": [],
    }


def _write_json(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    payload = {
        "schema_version": 1,
        "options": {"separator_mode": options.separator_mode, "custom_separator": options.custom_separator},
        "documents": [_record(item) for item in files],
    }
    with path.open("w", encoding="utf-8", newline="") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2, default=_json_value)
        stream.write("\n")


def _blocks(item: ConsolidationFile) -> list[Block]:
    return item.document.blocks if item.document is not None else [Block("paragraph", item.text)]


def _write_jsonl(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        for document_index, item in enumerate(files):
            record = _record(item)
            blocks = record.pop("blocks")
            for block_index, block in enumerate(blocks or [None]):
                row = {"schema_version": 1, "document_index": document_index,
                       "block_index": block_index if block is not None else None,
                       **record, "block": block}
                stream.write(json.dumps(row, ensure_ascii=False, default=_json_value) + "\n")


def _md(text: str) -> str:
    return re.sub(r"([\\\\`*_{}\[\]()#+.!|>~-])", r"\\\1", text).replace("&", "&amp;").replace("<", "&lt;")


def _write_md(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        for index, item in enumerate(files):
            if index:
                stream.write(_md(render_separator(options, item)))
            if options.separator_mode == "filename":
                stream.write(f"# {_md(item.name)}\n\n")
            for block in _blocks(item):
                text = _md(block.to_plain_text())
                if block.kind == "heading":
                    text = "#" * max(1, min(block.attrs.level or 1, 6)) + " " + text
                elif block.kind == "table" and block.attrs.table_data:
                    rows = block.attrs.table_data
                    width = max(map(len, rows))
                    lines = ["| " + " | ".join(_md(cell).replace("\n", "<br>") for cell in row + [""] * (width - len(row))) + " |" for row in rows]
                    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
                    text = "\n".join(lines)
                stream.write(text + "\n\n")


def _write_html(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        stream.write('<!doctype html><html><head><meta charset="utf-8"><title>Consolidation</title></head><body>\n')
        for index, item in enumerate(files):
            if index:
                stream.write("<pre>" + html.escape(render_separator(options, item)) + "</pre>\n")
            if options.separator_mode == "filename":
                stream.write("<h1>" + html.escape(item.name) + "</h1>\n")
            for block in _blocks(item):
                if block.kind == "table" and block.attrs.table_data:
                    stream.write("<table>" + "".join("<tr>" + "".join("<td>" + html.escape(cell).replace("\n", "<br>") + "</td>" for cell in row) + "</tr>" for row in block.attrs.table_data) + "</table>\n")
                else:
                    tag = f"h{max(1, min(block.attrs.level or 1, 6))}" if block.kind == "heading" else "pre"
                    stream.write(f"<{tag}>" + html.escape(block.text) + f"</{tag}>\n")
        stream.write("</body></html>\n")


def _write_docx(path: Path, files: Sequence[ConsolidationFile], options: ConsolidationOptions) -> None:
    from docx import Document as WordDocument

    word = WordDocument()
    for index, item in enumerate(files):
        if index:
            word.add_paragraph(render_separator(options, item))
        if options.separator_mode == "filename":
            word.add_heading(item.name, level=1)
        for block in _blocks(item):
            if block.kind == "heading":
                word.add_heading(block.text, level=max(1, min(block.attrs.level or 1, 9)))
            elif block.kind == "table" and block.attrs.table_data:
                rows = block.attrs.table_data
                width = max(map(len, rows))
                if width:
                    table = word.add_table(rows=len(rows), cols=width)
                    for row_index, row in enumerate(rows):
                        for column, text in enumerate(row):
                            table.cell(row_index, column).text = text
            elif block.kind == "page_break":
                word.add_page_break()
            else:
                word.add_paragraph(block.text)
    word.save(str(path))


def create_default_writer_registry() -> WriterRegistry:
    registry = WriterRegistry()
    registry.register("txt", _write_txt)
    registry.register("json", _write_json)
    registry.register("jsonl", _write_jsonl)
    registry.register("md", _write_md)
    registry.register("html", _write_html)
    registry.register("docx", _write_docx)
    return registry


def validate_destination(path: Path, sources: Sequence[Path], overwrite: bool = False) -> None:
    if not path.parent.is_dir():
        raise ValueError(f"Output directory does not exist: {path.parent}")
    if path.is_symlink():
        raise ValueError(f"Output must not be a symbolic link: {path}")
    for source in sources:
        if path.resolve() == source.resolve() or (
            path.exists() and source.exists() and path.samefile(source)
        ):
            raise ValueError(f"Output would overwrite a source: {path}")
    if path.exists() and (not overwrite or not path.is_file()):
        raise FileExistsError(f"Output already exists: {path}")


def write_output(
    path: Path,
    files: Sequence[ConsolidationFile],
    options: ConsolidationOptions,
    format: str | None = None,
    overwrite: bool = False,
    registry: WriterRegistry | None = None,
) -> None:
    writer = (registry or create_default_writer_registry()).get_writer(format or path.suffix)
    if any(item.error is not None for item in files):
        raise ValueError("Cannot export failed extraction results")
    sources = [item.path for item in files]
    validate_destination(path, sources, overwrite)
    descriptor, name = tempfile.mkstemp(prefix=".consolidar-", suffix=path.suffix, dir=path.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        writer(temporary, files, options)
        with temporary.open("rb+") as stream:
            os.fsync(stream.fileno())
        validate_destination(path, sources, overwrite)
        if overwrite:
            os.replace(temporary, path)
        else:
            os.link(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
