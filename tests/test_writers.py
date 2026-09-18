import json
import os

import pytest

from src.core.models import Block, BlockAttributes, Document, DocumentSource, ExtractionWarning
from src.core.writers import WriterRegistry
from pathlib import Path
from types import SimpleNamespace

from src.core.consolidator import build_output
from src.core.models import ConsolidationFile
from src.core.writers import write_output


def test_txt_matches_existing_rendering(tmp_path: Path) -> None:
    files = [ConsolidationFile(tmp_path / "a.txt", "a.txt", "Olá\r\nWorld"),
             ConsolidationFile(tmp_path / "b.txt", "b.txt", "second")]
    options = SimpleNamespace(separator_mode="filename", custom_separator="")
    output = tmp_path / "result.txt"
    write_output(output, files, options)
    assert output.read_bytes() == build_output(files, options).encode("utf-8")


@pytest.mark.parametrize("mode", ["blank", "dashes", "custom", "filename"])
def test_txt_separators(tmp_path: Path, mode: str) -> None:
    files = [ConsolidationFile(tmp_path / "a", "a", "A"), ConsolidationFile(tmp_path / "b", "b", "B")]
    options = SimpleNamespace(separator_mode=mode, custom_separator=r"\n{filename}\t")
    output = tmp_path / "result.txt"
    write_output(output, files, options)
    assert output.read_bytes() == build_output(files, options).encode()


def test_json_preserves_structure_and_provenance(tmp_path: Path) -> None:
    source = tmp_path / "input.txt"
    source.write_bytes(b"source")
    document = Document(DocumentSource.from_path(source))
    document.blocks = [Block("table", attrs=BlockAttributes(table_data=[["Olá", "2"]], page_number=1))]
    document.metadata.title = "Title"
    document.warnings = [ExtractionWarning("empty_page", "No text", 2)]
    item = ConsolidationFile(source, source.name, document.to_plain_text(), document=document)
    output = tmp_path / "result.json"
    write_output(output, [item], SimpleNamespace(separator_mode="blank", custom_separator=""))
    data = json.loads(output.read_text(encoding="utf-8"))
    assert data["schema_version"] == 1
    record = data["documents"][0]
    assert record["source"]["sha256"] == document.source.sha256
    assert record["source"]["mtime"] == document.source.mtime.isoformat()
    assert record["source"]["path"] == str(source.resolve())
    assert record["blocks"][0]["attrs"]["table_data"] == [["Olá", "2"]]
    assert record["metadata"]["title"] == "Title"
    assert record["warnings"][0]["page_number"] == 2


def test_source_and_existing_destination_protection(tmp_path: Path) -> None:
    source = tmp_path / "input.txt"
    source.write_bytes(b"original")
    item = ConsolidationFile(source, source.name, "new")
    options = SimpleNamespace(separator_mode="blank", custom_separator="")
    with pytest.raises(ValueError, match="source"):
        write_output(source, [item], options, overwrite=True)
    alias = tmp_path / "alias.txt"
    os.link(source, alias)
    with pytest.raises(ValueError, match="source"):
        write_output(alias, [item], options, overwrite=True)
    output = tmp_path / "out.txt"
    output.write_bytes(b"old")
    with pytest.raises(FileExistsError):
        write_output(output, [item], options)
    write_output(output, [item], options, overwrite=True)
    assert output.read_bytes() == b"new"
    assert source.read_bytes() == b"original"


def test_failed_writer_preserves_destination_and_cleans_temp(tmp_path: Path) -> None:
    output = tmp_path / "out.txt"
    output.write_bytes(b"old")
    def fail(path, files, options):
        path.write_bytes(b"partial")
        raise RuntimeError("writer failed")
    registry = WriterRegistry()
    registry.register("txt", fail)
    with pytest.raises(RuntimeError, match="writer failed"):
        write_output(output, [], SimpleNamespace(separator_mode="blank", custom_separator=""), overwrite=True, registry=registry)
    assert output.read_bytes() == b"old"
    assert list(tmp_path.iterdir()) == [output]


@pytest.mark.parametrize("format", ["txt", "json", "jsonl", "md", "html", "docx"])
def test_all_formats_preserve_structured_content(tmp_path: Path, format: str) -> None:
    source = tmp_path / "source.txt"
    source.write_bytes(b"source")
    document = Document(DocumentSource.from_path(source), blocks=[
        Block("heading", "Title", BlockAttributes(level=2)),
        Block("paragraph", "Body <&"),
        Block("table", attrs=BlockAttributes(table_data=[["Name", "Score"], ["Ana", "10"]])),
        Block("paragraph", "End"),
    ])
    item = ConsolidationFile(source, source.name, document.to_plain_text(), document=document)
    output = tmp_path / f"out.{format}"
    write_output(output, [item], SimpleNamespace(separator_mode="filename", custom_separator=""))
    if format == "docx":
        from docx import Document as WordDocument
        word = WordDocument(output)
        assert word.paragraphs[1].style.name == "Heading 2"
        assert word.tables[0].cell(1, 0).text == "Ana"
        assert word.paragraphs[-1].text == "End"
        assert [type(block).__name__ for block in word.iter_inner_content()] == [
            "Paragraph", "Paragraph", "Paragraph", "Table", "Paragraph",
        ]
    elif format == "json":
        record = json.loads(output.read_text(encoding="utf-8"))["documents"][0]
        assert record["blocks"][2]["attrs"]["table_data"][1] == ["Ana", "10"]
    elif format == "jsonl":
        records = [json.loads(line) for line in output.read_text(encoding="utf-8").splitlines()]
        assert [record["block_index"] for record in records] == [0, 1, 2, 3]
        assert all(record["source"]["sha256"] == document.source.sha256 for record in records)
        assert records[2]["block"]["kind"] == "table"
    else:
        text = output.read_text(encoding="utf-8")
        assert "Ana" in text and "End" in text
        if format == "html":
            assert "<h2>Title</h2>" in text
            assert "Body &lt;&amp;" in text
            assert "<td>Ana</td>" in text
            assert "&amp;lt;" not in text
        if format == "md":
            assert "## Title" in text
            assert "| Ana | 10 |" in text
            assert "Body &lt;&amp;" in text


def test_jsonl_retains_empty_documents(tmp_path: Path) -> None:
    source = tmp_path / "source.txt"
    source.touch()
    item = ConsolidationFile(source, source.name, document=Document(DocumentSource.from_path(source)))
    output = tmp_path / "out.jsonl"
    write_output(output, [item], SimpleNamespace(separator_mode="blank", custom_separator=""))
    record = json.loads(output.read_text(encoding="utf-8"))
    assert record["block"] is None
    assert record["block_index"] is None
    assert record["source"]["path"] == str(source)


def test_publication_does_not_overwrite_racing_destination(tmp_path: Path) -> None:
    output = tmp_path / "out.txt"
    def racing_writer(path, files, options):
        path.write_bytes(b"new")
        output.write_bytes(b"someone else's file")
    registry = WriterRegistry()
    registry.register("txt", racing_writer)
    with pytest.raises(FileExistsError):
        write_output(output, [], SimpleNamespace(separator_mode="blank", custom_separator=""), registry=registry)
    assert output.read_bytes() == b"someone else's file"
    assert list(tmp_path.iterdir()) == [output]


def test_unknown_format_and_failed_items_are_rejected(tmp_path: Path) -> None:
    options = SimpleNamespace(separator_mode="blank", custom_separator="")
    with pytest.raises(ValueError, match="Unsupported"):
        write_output(tmp_path / "out.unknown", [], options)
    with pytest.raises(ValueError, match="failed extraction"):
        write_output(tmp_path / "out.txt", [ConsolidationFile(Path("bad"), "bad", error="bad")], options)
