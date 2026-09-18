from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document as WordDocument

from src.core.consolidator import extract_document, extract_file
from src.extractors.docx import DocxExtractor
from src.extractors.registry import create_default_registry


def test_docx_preserves_heading_table_and_paragraph_order(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    word = WordDocument()
    word.add_heading("Report", level=1)
    word.add_paragraph("Before")
    table = word.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Name", "Score"], ["Ana", "10"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    word.add_heading("Details", level=2)
    word.add_paragraph("After")
    word.save(path)

    document = extract_document(path, create_default_registry())

    assert [block.kind for block in document.blocks] == [
        "heading", "paragraph", "table", "heading", "paragraph",
    ]
    assert document.blocks[0].attrs.level == 1
    assert document.blocks[3].attrs.level == 2
    assert document.blocks[2].attrs.table_data == [["Name", "Score"], ["Ana", "10"]]
    assert document.to_plain_text() == "Report\nBefore\nName\tScore\nAna\t10\nDetails\nAfter"
    assert DocxExtractor().extract(path) == document.to_plain_text()
    assert extract_file(path, create_default_registry()) == document.to_plain_text()


def test_docx_metadata(tmp_path: Path) -> None:
    path = tmp_path / "metadata.docx"
    word = WordDocument()
    created = datetime(2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
    word.core_properties.title = "Título"
    word.core_properties.author = "Author"
    word.core_properties.created = created
    word.core_properties.modified = created
    word.save(path)

    document = DocxExtractor().extract_structured(path)

    assert document.metadata.title == "Título"
    assert document.metadata.author == "Author"
    assert document.metadata.created == created
    assert document.metadata.modified == created
    assert document.metadata.page_count is None
    assert document.source.path == path.resolve()
    assert document.source.format == ".docx"


def test_docx_preserves_blank_paragraphs_and_cell_line_breaks(tmp_path: Path) -> None:
    path = tmp_path / "blank.docx"
    word = WordDocument()
    word.add_paragraph("First")
    word.add_paragraph("")
    word.add_paragraph("Last")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "line1\nline2"
    word.save(path)

    assert DocxExtractor().extract(path) == "First\n\nLast\nline1\nline2\t"


def test_empty_docx(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    WordDocument().save(path)

    document = DocxExtractor().extract_structured(path)

    assert document.blocks == []
    assert document.to_plain_text() == ""


def test_custom_style_based_on_heading(tmp_path: Path) -> None:
    from docx.enum.style import WD_STYLE_TYPE

    path = tmp_path / "styles.docx"
    word = WordDocument()
    style = word.styles.add_style("Report Section", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = word.styles["Heading 3"]
    word.add_paragraph("Section", style=style)
    word.save(path)

    block = DocxExtractor().extract_structured(path).blocks[0]

    assert block.kind == "heading"
    assert block.attrs.level == 3


def test_corrupt_docx_has_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")

    with pytest.raises(ValueError, match="DOCX inválido o corrupto"):
        DocxExtractor().extract_structured(path)
